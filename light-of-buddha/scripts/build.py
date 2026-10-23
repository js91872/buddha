#!/usr/bin/env python3
"""Build review editions; run from project root with python scripts/build.py."""
from pathlib import Path
import re, json, subprocess, zipfile, xml.etree.ElementTree as ET, html, os
ROOT=Path(__file__).resolve().parents[1]
os.chdir(ROOT)
REL=ROOT/'release'; REL.mkdir(exist_ok=True)
QA=ROOT/'publishing/quality-assurance'; QA.mkdir(parents=True,exist_ok=True)
TITLE='THE LIGHT OF BUDDHA'
SUB='30 Timeless Stories of Wisdom, Compassion & Inner Peace'
CREDIT='Compiled by Jaspal Singh'
PARTS={1:'PART I: The Prince Who Searched for Truth',8:'PART II: Lessons That Awaken the Heart',16:'PART III: Compassion That Changes Lives',24:'PART IV: Wisdom for Every Day'}

def norm(s):
 s=re.sub(r'\]\((?:\.\./)+artwork/','](artwork/',s)
 s=re.sub(r'\]\((?:\.\./)+research/','](research/',s)
 return s

def run(args, name):
 p=subprocess.run(args,text=True,capture_output=True)
 (QA/name).write_text('$ '+' '.join(map(str,args))+'\n'+p.stdout+p.stderr+'\nExit: '+str(p.returncode)+'\n')
 if p.returncode: raise RuntimeError(name+' failed; see log')

stories=sorted((ROOT/'manuscript/stories').glob('[0-9][0-9]*.md'))
import sys
if '--final' in sys.argv and [int(p.name[:2]) for p in stories] != list(range(1,31)):raise SystemExit('Final build requires exactly stories01–30')
front=sorted((ROOT/'manuscript/front-matter').glob('*.md'))
back=sorted((ROOT/'manuscript/back-matter').glob('*.md'))
verses=sorted((ROOT/'manuscript/verses-and-chants').glob('[0-9][0-9]-interlude.md'))
verse_map={int(e['after_story']):int(e['interlude']) for e in json.loads((ROOT/'manuscript/verses-and-chants/insertion-map.json').read_text())} if (ROOT/'manuscript/verses-and-chants/insertion-map.json').exists() else {}
if '--final' in sys.argv:
 if len(front)<5 or len(back)<7 or len(verses)!=13:raise SystemExit('Final build requires front/back matter and thirteen interludes')
 for n in range(1,31):
  if not (ROOT/f'artwork/story-illustrations/{n:02}.jpg').is_file():raise SystemExit(f'Missing story illustration {n:02}')
 if not (REL/'kindle-cover.jpg').is_file():raise SystemExit('Missing approved Kindle marketing cover')
 if CREDIT not in '\n'.join(p.read_text() for p in front):raise SystemExit('Approved contributor display credit missing from front matter')
for n in range(1,31):
 if not (ROOT/f'publishing/kindle/images/{n:02}.jpg').is_file():raise SystemExit(f'Missing optimized Kindle image {n:02}')
chunks=[p.read_text() for p in front]
if not front:
 chunks=[f'# {TITLE}\n\n{SUB}\n\nVOLUME ONE\n\n{CREDIT}\n']
for p in stories:
 n=int(p.name[:2])
 if n in PARTS: chunks.append('# '+PARTS[n]+'\n')
 text=norm(p.read_text())
 # Every numbered story gets one accessible source illustration if writer omitted it.
 if not re.search(r'!\[',text):
  alt=(ROOT/f'artwork/alt-text/{n:02}.txt').read_text().strip()
  lines=text.splitlines(); lines.insert(1,f'\n![{alt}](artwork/story-illustrations/{n:02}.jpg)\n'); text='\n'.join(lines)
 chunks.append(text)
 # Insert the selected interlude according to its explicit story placement map.
 if n in verse_map: chunks.append(verses[verse_map[n]-1].read_text())
chunks.extend(p.read_text() for p in back)
combined='\n\n'.join(chunks)
(ROOT/'manuscript/complete-manuscript.md').write_text(combined)
kindle_source=combined
for image_number in range(1,31):
 kindle_source=kindle_source.replace(f'artwork/story-illustrations/{image_number:02}.jpg',f'publishing/kindle/images/{image_number:02}.jpg')
(ROOT/'publishing/kindle/complete-manuscript-kindle.md').write_text(kindle_source)
css='''body{font-family:serif;line-height:1.45;margin:0 4%;orphans:3;widows:3} h1,h2,h3{line-height:1.2;text-align:left;break-after:avoid;page-break-after:avoid} h1{margin:1.25em 0 .65em;break-before:avoid-page;page-break-before:avoid} h2,h3{margin:1.1em 0 .5em} p{text-indent:0;margin:0 0 .8em} img{display:block;margin:1em auto;max-width:100%;width:auto;height:auto} figure{text-align:center;margin:1.2em 0} figcaption{font-size:.8em} a{overflow-wrap:anywhere}'''
(ROOT/'publishing/kindle/style.css').write_text(css)
meta='''<?xml version="1.0" encoding="UTF-8"?>
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:opf="http://www.idpf.org/2007/opf"><dc:identifier>urn:uuid:6332597e-c11d-5d62-a639-d50443bea97b</dc:identifier><dc:title>THE LIGHT OF BUDDHA</dc:title><dc:language>en</dc:language><dc:contributor id="compiler">Jaspal Singh</dc:contributor><meta refines="#compiler" property="role" scheme="marc:relators">com</meta><dc:description>30 Timeless Stories of Wisdom, Compassion &amp; Inner Peace. Volume One. Compiled by Jaspal Singh.</dc:description></metadata>'''
(ROOT/'publishing/metadata/epub-metadata.xml').write_text(meta)
args=['pandoc','manuscript/complete-manuscript.md','--from=markdown-implicit_figures','--standalone','--toc','--toc-depth=1','--metadata=lang:en','--resource-path=.:manuscript']
cover=ROOT/'release/kindle-cover.jpg'
epub=args.copy(); epub[1]='publishing/kindle/complete-manuscript-kindle.md'
epub+=['--metadata=title:'+TITLE,'--metadata=subtitle:'+SUB,'--to=epub3','--split-level=6','--epub-title-page=false','--epub-metadata=publishing/metadata/epub-metadata.xml','--css=publishing/kindle/style.css']
if cover.exists(): epub+=['--epub-cover-image='+str(cover)]
run(epub+['-o',str(REL/'the-light-of-buddha-volume-1.epub')],'epub-build.log')
# Explicitly install Dublin Core metadata; Pandoc versions differ in XML-wrapper handling.
from lxml import etree as LX
from tempfile import NamedTemporaryFile
bookpath=REL/'the-light-of-buddha-volume-1.epub'
with zipfile.ZipFile(bookpath) as zin:
 entries=[(info,zin.read(info.filename)) for info in zin.infolist()]
for index,(info,data) in enumerate(entries):
 if not info.filename.endswith('.opf'):continue
 tree=LX.fromstring(data);ns={'o':'http://www.idpf.org/2007/opf','d':'http://purl.org/dc/elements/1.1/'}
 md=tree.find('o:metadata',ns)
 for el in list(md):
  if el.tag in ['{'+ns['d']+'}title','{'+ns['d']+'}contributor','{'+ns['d']+'}description'] or (el.tag=='{'+ns['o']+'}meta' and el.get('refines') in ['#main-title','#subtitle','#compiler']):md.remove(el)
 def dc(tag,value,**attrs):
  el=LX.SubElement(md,'{'+ns['d']+'}'+tag,**attrs);el.text=value;return el
 def ref(target,prop,value,**attrs):
  el=LX.SubElement(md,'{'+ns['o']+'}meta',refines='#'+target,property=prop,**attrs);el.text=value
 dc('title',TITLE,id='main-title');ref('main-title','title-type','main')
 dc('title',SUB,id='subtitle');ref('subtitle','title-type','subtitle')
 dc('contributor','Jaspal Singh',id='compiler');ref('compiler','role','com',scheme='marc:relators')
 dc('description',SUB+'. Volume One. '+CREDIT+'.')
 collection=LX.SubElement(md,'{'+ns['o']+'}meta',property='belongs-to-collection',id='series');collection.text='The Light of Buddha'
 ref('series','collection-type','series');ref('series','group-position','1')
 identifier=md.find('d:identifier',ns)
 identifier.text='urn:uuid:6332597e-c11d-5d62-a639-d50443bea97b'
 entries[index]=(info,LX.tostring(tree,xml_declaration=True,encoding='UTF-8'))
for index,(info,data) in enumerate(entries):
 if not info.filename.endswith('.ncx'):continue
 tree=LX.fromstring(data);ns={'n':'http://www.daisy.org/z3986/2005/ncx/'}
 for el in tree.xpath('//n:meta[@name="dtb:uid"]',namespaces=ns):el.set('content','urn:uuid:6332597e-c11d-5d62-a639-d50443bea97b')
 for el in tree.xpath('//n:docTitle/n:text',namespaces=ns):el.text=TITLE
 entries[index]=(info,LX.tostring(tree,xml_declaration=True,encoding='UTF-8'))
for index,(info,data) in enumerate(entries):
 if not info.filename.endswith('cover.xhtml'):continue
 tree=LX.fromstring(data);svgns='http://www.w3.org/2000/svg'
 for svg in tree.iter('{'+svgns+'}svg'):
  svg.set('role','img');svg.set('aria-labelledby','cover-title cover-description')
  title=LX.Element('{'+svgns+'}title',id='cover-title');title.text=TITLE
  desc=LX.Element('{'+svgns+'}desc',id='cover-description');desc.text='Serene Buddha meditating beneath a Bodhi tree. '+SUB+'. VOLUME ONE. '+CREDIT+'.'
  svg.insert(0,desc);svg.insert(0,title)
 entries[index]=(info,LX.tostring(tree,xml_declaration=True,encoding='UTF-8'))
with zipfile.ZipFile(bookpath,'w') as zout:
 for info,data in entries:zout.writestr(info,data)

run(args+['--to=docx','-o',str(REL/'the-light-of-buddha-volume-1.docx')],'docx-build.log')
from docx import Document as WordDocument
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
wd=WordDocument(REL/'the-light-of-buddha-volume-1.docx')
for section in wd.sections:
 section.page_width=Inches(6);section.page_height=Inches(9)
 section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(.65)
wordstyles={sty.name:sty for sty in wd.styles}
for sty in ['Normal','Body Text','First Paragraph']:
 if sty in wordstyles:
  wordstyles[sty].font.name='Noto Serif';wordstyles[sty].font.size=Pt(10.5)
  wordstyles[sty].paragraph_format.line_spacing=1.3
wordstyles['Heading 1'].font.name='Noto Serif';wordstyles['Heading 1'].font.size=Pt(19)
wordstyles['Heading 1'].paragraph_format.page_break_before=True
for paragraph in wd.paragraphs:
 if paragraph.text in [TITLE,SUB,CREDIT,'VOLUME ONE','Volume 1','The Light of Buddha — Volume 1']:
  paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
for shape in wd.inline_shapes:
 ratio=shape.height/shape.width
 shape.width=Inches(3.74);shape.height=int(Inches(3.74)*ratio)
wd.core_properties.title=TITLE;wd.core_properties.subject=SUB;wd.core_properties.author=CREDIT
wd.save(REL/'the-light-of-buddha-volume-1.docx')


# Print edition: six by nine inches, generous no-bleed margins, embedded Unicode fonts.
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,Image,KeepTogether
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.utils import ImageReader
from PIL import Image as PILImage
fontdir=Path('/usr/share/fonts/truetype/noto')
for key,file in [('Book','NotoSerif-Regular.ttf'),('BookB','NotoSerif-Bold.ttf'),('BookI','NotoSerif-Italic.ttf')]: pdfmetrics.registerFont(TTFont(key,str(fontdir/file)))
pdfmetrics.registerFontFamily('Book',normal='Book',bold='BookB',italic='BookI',boldItalic='BookB')
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='BookBody',fontName='Book',fontSize=10.4,leading=14.5,spaceAfter=4.5,allowWidows=0,allowOrphans=0))
styles.add(ParagraphStyle(name='BookH1',fontName='BookB',fontSize=19,leading=25,spaceAfter=18,keepWithNext=True))
styles.add(ParagraphStyle(name='BookH2',fontName='BookB',fontSize=11,leading=15,spaceBefore=10,spaceAfter=5,keepWithNext=True))
styles.add(ParagraphStyle(name='BookSmall',fontName='Book',fontSize=8.5,leading=11,spaceAfter=5))
styles.add(ParagraphStyle(name='BookSourceH',fontName='BookB',fontSize=10,leading=12,spaceBefore=6,spaceAfter=3,keepWithNext=True))
styles.add(ParagraphStyle(name='BookCopyright',fontName='Book',fontSize=9,leading=12,spaceAfter=5))
styles.add(ParagraphStyle(name='BookClosing',fontName='Book',fontSize=10.4,leading=14,spaceAfter=4))
styles.add(ParagraphStyle(name='BookClosingH',fontName='BookB',fontSize=17,leading=22,spaceAfter=14,keepWithNext=True))

def inline(s):
 s=html.escape(s)
 s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',lambda m:'<link href="'+html.escape(m[2],quote=True)+'" color="#285b65">'+m[1]+'</link>',s)
 s=re.sub(r'\*\*(.+?)\*\*',r'<b>\1</b>',s)
 s=re.sub(r'(?<!\*)\*([^*]+)\*',r'<i>\1</i>',s)
 s=re.sub(r'`([^`]+)`',r'\1',s)
 return s
styles.add(ParagraphStyle(name='BookTitle',fontName='BookB',fontSize=25,leading=34,alignment=TA_CENTER,spaceAfter=22))
styles.add(ParagraphStyle(name='BookCredit',fontName='Book',fontSize=12,leading=18,alignment=TA_CENTER,spaceAfter=14))
class BookDoc(SimpleDocTemplate):
 def __init__(self,*args,**kwargs):
  super().__init__(*args,**kwargs); self.heading_counter=0
 def beforeDocument(self): self.heading_counter=0
 def afterFlowable(self,obj):
  if isinstance(obj,Paragraph) and obj.style.name=='BookH1':
   self.heading_counter+=1
   title=obj.getPlainText(); key='heading-%03d'%self.heading_counter
   self.canv.bookmarkPage(key); self.canv.addOutlineEntry(title,key,level=0,closed=False)
   self.notify('TOCEntry',(0,title,self.page,key))
toc=TableOfContents()
toc.levelStyles=[ParagraphStyle(name='TOCLine',fontName='Book',fontSize=9,leading=13,spaceBefore=0,leftIndent=0,rightIndent=20)]
toc.tableStyle.add('TOPPADDING',(0,0),(-1,-1),2)
toc.tableStyle.add('BOTTOMPADDING',(0,0),(-1,-1),2)
flow=[]
for chunk_index,chunk in enumerate(chunks):
 in_source_note=False
 tail_start=None
 is_closing=chunk.casefold().startswith('# closing reflection')
 if chunk_index==len(front):
  if flow: flow.append(PageBreak())
  flow.extend([Paragraph('Contents',styles['BookTitle']),toc])
 if flow and (chunk_index < len(front) or chunk_index == len(front) or chunk.lstrip().startswith('# PART ')): flow.append(PageBreak())
 for block in re.split(r'\n\s*\n',norm(chunk)):
  block=block.strip()
  if not block: continue
  if block.startswith('!['):
   m=re.search(r'!\[.*?\]\(([^)]+)\)',block,re.S)
   if m and (ROOT/m[1]).exists():
    im=ROOT/m[1]; w,h=PILImage.open(im).size
    width=3.74*72; height=width*h/w
    if height>5.4*72: width*=5.4*72/height; height=5.4*72
    obj=Image(str(im),width=width,height=height); obj.hAlign='CENTER'; flow.extend([obj,Spacer(1,12)])
   continue
  if block.startswith('#'):
   lines=block.splitlines(); heading=lines.pop(0); level=len(heading)-len(heading.lstrip('#'))
   label=heading.lstrip('#').strip()
   if label==TITLE:
    flow.extend([Spacer(1,80),Paragraph(inline(label),styles['BookTitle'])])
   else:
    if label.casefold().startswith('wisdom to remember'):tail_start=len(flow)
    if label.casefold().startswith('source note'):
     if tail_start is not None:
      for tail_flow in flow[tail_start:]:tail_flow.keepWithNext=True
     in_source_note=True; heading_style='BookSourceH'
    else:heading_style='BookClosingH' if is_closing and level==1 else 'BookH1' if level==1 else 'BookH2'
    flow.append(Paragraph(inline(label),styles[heading_style]))
   if not lines: continue
   block='\n'.join(lines)
  if block.startswith('---') and block.count('---')>=2: continue
  block=re.sub(r'^>\s?','',block,flags=re.M)
  block=re.sub(r'^[-*] ','• ',block,flags=re.M)
  block='\n'.join(' · '.join(c.strip() for c in line.strip('|').split('|')) if line.startswith('|') and not re.match(r'^\|[ :|\-]+$',line) else '' if re.match(r'^\|[ :|\-]+$',line) else line for line in block.splitlines())
  if not block.strip(): continue
  style='BookSmall' if in_source_note or block.startswith(('Source','**Source','Canonical source')) else 'BookBody'
  if chunk_index < len(front) and 'Publication Information' in chunk:style='BookCopyright'
  if is_closing:style='BookClosing'
  if chunk_index < len(front) and ((chunk_index in [0,1]) or block==CREDIT):style='BookCredit'
  flow.append(Paragraph(inline(block).replace('\n','<br/>'),styles[style]))
def footer(c,d):
 c.setFont('Book',8); c.setFillColorRGB(.3,.35,.35); c.drawCentredString(3*72,.38*72,str(d.page))
 c.setTitle(TITLE); c.setAuthor(CREDIT); c.setSubject(SUB)
pdf=REL/'the-light-of-buddha-volume-1-print.pdf'
doc=BookDoc(str(pdf),pagesize=(6*72,9*72),rightMargin=.65*72,leftMargin=.65*72,topMargin=.65*72,bottomMargin=.65*72,title=TITLE,author=CREDIT)
doc.multiBuild(flow,onFirstPage=footer,onLaterPages=footer,canvasmaker=lambda *a,**k:canvas.Canvas(*a,**dict(k,initialFontName='Book',initialFontSize=10)))
# Structural checks are complementary to EPUBCheck, not a replacement.
report={'stories':len(stories),'verses':len(verses),'front_matter_files':len(front),'back_matter_files':len(back),'cover_present':cover.exists(),'epubcheck':'not run by build.py; see dedicated log','kindle_previewer':'unavailable in this environment'}
with zipfile.ZipFile(REL/'the-light-of-buddha-volume-1.epub') as z:
 bad=[]; images=0; missing_alt=[]
 for name in z.namelist():
  if name.endswith(('.xhtml','.opf','.ncx','.xml')):
   try: tree=ET.fromstring(z.read(name))
   except ET.ParseError as e: bad.append([name,str(e)]); continue
   for img in tree.iter('{http://www.w3.org/1999/xhtml}img'):
    images+=1
    if not img.get('alt'): missing_alt.append(name)
 report.update(epub_xml_errors=bad,epub_images=images,epub_missing_alt=missing_alt,epub_has_navigation=any(n.endswith('nav.xhtml') for n in z.namelist()),epub_mimetype=z.read('mimetype').decode())
import fitz,shutil
archive=QA/'interim-layouts';archive.mkdir(exist_ok=True)
for old in list(QA.glob('print-page-*.png'))+list(QA.glob('inspect-*.png')):shutil.move(str(old),str(archive/old.name))
with fitz.open(pdf) as d:
 report['pdf_pages']=len(d); report['pdf_page_size_points']=list(d[0].rect)[2:]
 report['pdf_empty_pages']=[i+1 for i,p in enumerate(d) if not p.get_text().strip() and not p.get_images()]
 d[0].get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(str(QA/'print-page-001.png'))
 for i in sorted(set([min(x-1,len(d)-1) for x in [2,3,10,30,60,65,124,len(d)]])):d[i].get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(str(QA/f'print-page-{i+1:03}.png'))
(QA/'build-checks.json').write_text(json.dumps(report,indent=2))
print(json.dumps(report,indent=2))
